"""Checks the recalculated Excel workbook equals the Python analysis. Run from the repo root."""
import sys; sys.path.insert(0, "src")
from openpyxl import load_workbook
import analysis as A
from themes import GROUPS, theme_names

brands, _ = A.load()
pos, cta, eng = (A.theme_matrix(brands, g) for g in ("positioning", "cta", "engagement"))
m = A.metrics(brands); k = A.key_numbers(brands, pos, cta, eng, m)
wb = load_workbook("excel/gtm_workbook.xlsx", data_only=True)
bad = 0
def chk(label, a, b):
    global bad
    if a != b and not (isinstance(a, (int, float)) and isinstance(b, (int, float)) and abs(a - b) < 1e-9):
        print("MISMATCH", label, a, b); bad += 1

mt = wb["Metrics"]
for i, b in enumerate(brands.brand):
    r = 4 + i; row = m[m.brand == b].iloc[0]; assert mt.cell(r, 1).value == b
    for col, key in [(2, "industries_n"), (3, "engagement_n"), (4, "cta_n"), (5, "partner_n"), (6, "compliance_n"), (7, "content_score"), (8, "publishes_rate")]:
        chk(f"{b}.{key}", mt.cell(r, col).value, row[key])
for sheet, group, df in [("Positioning", "positioning", pos), ("CTA Funnel", "cta", cta), ("Engagement", "engagement", eng)]:
    ws = wb[sheet]; names = theme_names(group)
    first = 6 + len(GROUPS[group]["keywords"]) + 2
    hdr = {ws.cell(4, c).value: c for c in range(5, 5 + len(brands))}
    for i, t in enumerate(names):
        assert ws.cell(first + i, 1).value == t, (sheet, t)
        for b, c in hdr.items():
            chk(f"{sheet}.{b}.{t}", ws.cell(first + i, c).value, int(df[df.brand == b][t].iloc[0]))
gp = wb["GTM Plan"]; xl = {gp.cell(r, 1).value: gp.cell(r, 2).value for r in range(5, 22)}
pairs = {"Brands compared": "n_brands", "Headlines AI-led": "ai_led", "Headlines naming a vertical": "vertical_headline",
         "Brands publishing a rate": "publish_rate", "Call / consult CTA": "call_cta", "Quote / estimate CTA": "quote_cta",
         "NDA-first CTA": "nda_cta", "Show-proof CTA": "proof_cta", "Hire-talent CTA": "hire_cta", "Case studies: Yes": "case_yes",
         "Whitepapers: Yes": "whitepaper_yes", "Podcast/newsletter: Yes": "podcast_yes", "Cite a Clutch rating": "clutch_cited",
         "Name a dedicated-team model": "dedicated_team", "Name a fixed-price model": "fixed_price",
         "Median industries listed": "median_industries", "Median content score": "median_content"}
for lab, key in pairs.items(): chk(lab, xl[lab], k[key])
plan = A.gtm_plan(k, m, brands)
for i, p in enumerate(plan):
    chk(f"plan{i} evidence", gp.cell(5 + len(pairs) + 2 + i, 2).value, p["evidence"])
print("parity mismatches:", bad); sys.exit(bad > 0)
