"""Builds excel/gtm_workbook.xlsx. Data sheets hold sourced inputs (blue); every metric is a live formula
(black). Keyword detection mirrors src/themes.py exactly."""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from themes import GROUPS, theme_names, PUNCT

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw"
BLUE, BOLD = Font(color="0000FF"), Font(bold=True)
HDR, HF = PatternFill("solid", fgColor="1F3864"), Font(bold=True, color="FFFFFF")
brands = pd.read_csv(RAW / "brands.csv"); gaps = pd.read_csv(RAW / "data_gaps.csv")
NB = len(brands); BR = list(brands.brand)
COL = {c: L(i + 1) for i, c in enumerate(brands.columns)}   # column letter by name in Brands sheet


def rng(c): return f"Brands!${COL[c]}$2:${COL[c]}${NB+1}"
def cell(c, i): return f"Brands!${COL[c]}${i+2}"


def dump(ws, df, widths=None):
    for j, c in enumerate(df.columns, 1):
        x = ws.cell(1, j, c); x.font, x.fill = HF, HDR
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            v = None if pd.isna(v) else (v.item() if hasattr(v, "item") else v)
            x = ws.cell(i, j, v); x.font = BLUE; x.alignment = Alignment(wrap_text=True, vertical="top")
    for j in range(1, len(df.columns) + 1): ws.column_dimensions[L(j)].width = (widths or {}).get(j, 18)
    ws.freeze_panes = "B2"


def header(ws, r, vals, c0=1):
    for j, v in enumerate(vals, c0):
        x = ws.cell(r, j, v); x.font, x.fill = HF, HDR; x.alignment = Alignment(wrap_text=True, vertical="center")


wb = Workbook()
ab = wb.active; ab.title = "About"
for i, t in enumerate([
    "Market Intelligence & GTM Dashboard: software engineering services firms",
    "Public data from each firm's own homepage, accessed 2026-10-12. See Sources & Gaps.",
    "Blue text = sourced input. Black text = formula. Change an input and every metric updates.",
    "Wingman Partners itself is NOT in the dataset (its site could not be read). Recommendations are framed for a firm of its type.",
    "Years, team sizes, ratings and badges are the firms' own claims, not verified data.",
    "'Not found' = not seen in the homepage text, not proven absent.",
    "Not collected: traffic, SEO, ads, LinkedIn activity, revenue, win rates. This is a public-homepage benchmark, not market research."], 1):
    ab.cell(i, 1, t).font = Font(bold=(i == 1), size=14 if i == 1 else 11)
ab.column_dimensions["A"].width = 130

wsB = wb.create_sheet("Brands")
dump(wsB, brands, {2: 28, 3: 36, 4: 44, 6: 40, 7: 40, 8: 30, 11: 40, 14: 40, 15: 24, 16: 36, 17: 30, 26: 50, 27: 30})
wsG = wb.create_sheet("Sources & Gaps"); dump(wsG, gaps, {1: 22, 2: 70, 3: 70})


def kw_sheet(name, group):
    g = GROUPS[group]; ws = wb.create_sheet(name)
    ws["A1"] = g["title"]; ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Text is lower-cased, punctuation replaced by spaces, padded with spaces. 'word' = whole word, 'prefix' = word starts with keyword. Edit the blue keywords to re-run."
    header(ws, 4, ["Theme", "Keyword", "Match", "Search needle"] + BR)
    ws.cell(5, 1, "Normalised text").font = BOLD
    for j in range(NB):
        src = "&\" \"&".join(cell(c, j) for c in g["columns"])
        expr = f"LOWER({src})"
        for ch in PUNCT:
            expr = f'SUBSTITUTE({expr},"{ch if ch != chr(34) else chr(34)*2}"," ")'
        ws.cell(5, 5 + j, f'=" "&{expr}&" "')
    K0 = 6
    for i, (t, kw, m) in enumerate(g["keywords"]):
        r = K0 + i
        for c, v in enumerate([t, kw, m], 1): ws.cell(r, c, v).font = BLUE
        ws.cell(r, 4, f'=" "&B{r}&IF(C{r}="word"," ","")')
        for j in range(NB):
            cl = L(5 + j)
            ws.cell(r, 5 + j, f'=IF(B{r}="",0,IF(ISNUMBER(FIND($D{r},{cl}$5)),1,0))')
    K1 = K0 + len(g["keywords"]) - 1; S0 = K1 + 3
    header(ws, S0 - 1, ["Theme summary (1 = present)", "", "", "Brands using"] + BR)
    names = theme_names(group)
    for i, t in enumerate(names):
        r = S0 + i; ws.cell(r, 1, t)
        for j in range(NB):
            cl = L(5 + j)
            ws.cell(r, 5 + j, f"=IF(SUMIF($A${K0}:$A${K1},$A{r},{cl}${K0}:{cl}${K1})>0,1,0)")
        ws.cell(r, 4, f"=SUM(E{r}:{L(4+NB)}{r})")
    ws.column_dimensions["A"].width = 30
    for c in "BCD": ws.column_dimensions[c].width = 16
    for j in range(NB): ws.column_dimensions[L(5 + j)].width = 14
    return ws, S0, S0 + len(names) - 1, names


posS, pS0, pS1, pNames = kw_sheet("Positioning", "positioning")
ctaS, cS0, cS1, cNames = kw_sheet("CTA Funnel", "cta")
engS, eS0, eS1, eNames = kw_sheet("Engagement", "engagement")

# ---------- Metrics
mt = wb.create_sheet("Metrics")
mt["A1"] = "Per-brand metrics (formulas on Brands sheet)"; mt["A1"].font = Font(bold=True, size=13)
header(mt, 3, ["Brand", "Industries listed", "Engagement models named", "CTAs listed", "Partner badges", "Compliance items listed",
               "Content engine score (0-6)", "Publishes a rate", "Years claimed", "Team claimed"])


def cnt(ref): return f'IF({ref}="",0,LEN({ref})-LEN(SUBSTITUTE({ref},";",""))+1)'


for i, b in enumerate(BR):
    r = 4 + i
    mt.cell(r, 1, f"={cell('brand', i)}")
    mt.cell(r, 2, "=" + cnt(cell("industries", i))); mt.cell(r, 3, "=" + cnt(cell("engagement_models", i)))
    mt.cell(r, 4, "=" + cnt(cell("primary_ctas", i))); mt.cell(r, 5, "=" + cnt(cell("partner_badges", i)))
    mt.cell(r, 6, "=" + cnt(cell("compliance_listed", i)))
    mt.cell(r, 7, "=" + "+".join(f'IF({cell(c, i)}="Yes",1,IF({cell(c, i)}="Partial",0.5,0))'
                                  for c in ["case_studies", "blog", "whitepapers", "testimonials", "client_logos", "podcast_newsletter"]))
    mt.cell(r, 8, f'=IF({cell("rate_low_usd_hr", i)}="","No","Yes")')
    mt.cell(r, 9, f'=IF({cell("years_claimed", i)}="","",{cell("years_claimed", i)})')
    mt.cell(r, 10, f'=IF({cell("team_claimed", i)}="","",{cell("team_claimed", i)})')
M1 = 3 + NB
mt.cell(M1 + 2, 1, "Median industries listed").font = BOLD; mt.cell(M1 + 2, 2, f"=MEDIAN(B4:B{M1})")
mt.cell(M1 + 3, 1, "Median content engine score").font = BOLD; mt.cell(M1 + 3, 2, f"=MEDIAN(G4:G{M1})")
mt.cell(M1 + 4, 1, "Median partner badges").font = BOLD; mt.cell(M1 + 4, 2, f"=MEDIAN(E4:E{M1})")
mt.column_dimensions["A"].width = 30
for c in "BCDEFGHIJ": mt.column_dimensions[c].width = 16
MED = {"ind": M1 + 2, "content": M1 + 3, "partner": M1 + 4}

# ---------- GTM Plan
gp = wb.create_sheet("GTM Plan")
gp["A1"] = "GTM recommendations derived from the data (hypotheses to test, not facts)"; gp["A1"].font = Font(bold=True, size=13)
gp["A2"] = "Wingman Partners' own site could not be read; these are framed for a software-services firm of its type."
header(gp, 4, ["Key number (live)", "Value"])
def tcol(sheet, S0, names, theme, j=None):
    r = S0 + names.index(theme); return f"'{sheet}'!$D${r}"
KN = [
    ("Brands compared", f"=COUNTA({rng('brand')})"),
    ("Headlines AI-led", f"={tcol('Positioning', pS0, pNames, 'AI-led')}"),
    ("Headlines naming a vertical", f"={tcol('Positioning', pS0, pNames, 'Vertical focus')}"),
    ("Brands publishing a rate", f'=COUNTIF(Metrics!$H$4:$H${M1},"Yes")'),
    ("Call / consult CTA", f"={tcol('CTA Funnel', cS0, cNames, 'Book a call / consult')}"),
    ("Quote / estimate CTA", f"={tcol('CTA Funnel', cS0, cNames, 'Get a quote / estimate')}"),
    ("NDA-first CTA", f"={tcol('CTA Funnel', cS0, cNames, 'NDA-first trust')}"),
    ("Show-proof CTA", f"={tcol('CTA Funnel', cS0, cNames, 'Show proof')}"),
    ("Hire-talent CTA", f"={tcol('CTA Funnel', cS0, cNames, 'Hire talent')}"),
    ("Case studies: Yes", f'=COUNTIF({rng("case_studies")},"Yes")'),
    ("Whitepapers: Yes", f'=COUNTIF({rng("whitepapers")},"Yes")'),
    ("Podcast/newsletter: Yes", f'=COUNTIF({rng("podcast_newsletter")},"Yes")'),
    ("Cite a Clutch rating", f'=COUNTIF({rng("rating_claim")},"*Clutch*")'),
    ("Name a dedicated-team model", f"={tcol('Engagement', eS0, eNames, 'Dedicated team')}"),
    ("Name a fixed-price model", f"={tcol('Engagement', eS0, eNames, 'Fixed price')}"),
    ("Median industries listed", f"=Metrics!B{MED['ind']}"),
    ("Median content score", f"=Metrics!B{MED['content']}"),
]
KROW = {}
for i, (lab, f) in enumerate(KN):
    r = 5 + i; gp.cell(r, 1, lab); gp.cell(r, 2, f); KROW[lab] = r
def kv(lab): return f"$B${KROW[lab]}"
PL0 = 5 + len(KN) + 2
header(gp, PL0 - 1, ["Insight", "Evidence (live)", "Recommendation", "Test to run", "Metric", "Caveat"])
n = kv("Brands compared")
plan_rows = [
 ("Positioning: AI is table stakes, a specific outcome or vertical is not",
  f'={kv("Headlines AI-led")}&"/"&{n}&" headlines lead with AI; only "&{kv("Headlines naming a vertical")}&"/"&{n}&" name a vertical."',
  "Do not make 'AI' the whole headline. Lead with the problem solved for one named buyer, with AI as the how.",
  "A/B two homepage headlines: AI-led vs. outcome/vertical-led.", "Click-through on the main CTA; booked calls per 100 visitors.",
  "Homepage headline + subheadline only; says nothing about performance."),
 ("Pricing: almost nobody shows a number",
  f'={kv("Brands publishing a rate")}&"/"&{n}&" brands publish any rate; the rest hide pricing behind a quote."',
  "A 'typical engagement' guide (ranges by team size or scope) can act as trust-builder and lead magnet.",
  "Gated vs. ungated pricing guide.", "Pricing-page visit to qualified-lead rate; sales-accepted lead rate.",
  "The one published rate belongs to a healthcare-niche vendor; one data point."),
 ("Lead capture: everyone says 'talk to us', few offer a smaller first step",
  f'={kv("Call / consult CTA")}&"/"&{n}&" push a call/consult CTA and "&{kv("Quote / estimate CTA")}&"/"&{n}&" an estimate; only "&{kv("NDA-first CTA")}&"/"&{n}&" offer an NDA-first step."',
  "Add a lower-commitment step beside the call CTA (estimate tool, scoping checklist, NDA-first intake).",
  "Call CTA only vs. call CTA + estimate/checklist CTA.", "Total leads, and the share that become sales-qualified.",
  "CTA text read from the homepage; placement and form design not captured."),
 ("Content: case studies are the baseline, deeper assets are rare",
  f'={kv("Case studies: Yes")}&"/"&{n}&" show case studies, but only "&{kv("Whitepapers: Yes")}&"/"&{n}&" show whitepapers and "&{kv("Podcast/newsletter: Yes")}&"/"&{n}&" a podcast or newsletter."',
  "Treat case studies as required. Differentiate with a technical point of view for one niche (gated whitepaper, newsletter).",
  "Gate one whitepaper behind an email form; promote it on LinkedIn to the target buyer.", "Downloads, cost per lead, lead-to-call rate.",
  "'Not found' means not seen on the homepage, not proven absent."),
 ("Focus: breadth is the norm, a niche is rare",
  f'="Industries listed per firm range from "&MIN(Metrics!$B$4:$B${M1})&" to "&MAX(Metrics!$B$4:$B${M1})&" (median "&{kv("Median industries listed")}&"); "&{kv("Headlines naming a vertical")}&"/"&{n}&" headline a single vertical."',
  "Pick one or two verticals as the ideal customer profile and build proof for them.",
  "Run a vertical-specific landing page against the generic one.", "Conversion by landing page; lead quality by segment.",
  "Wingman's own verticals are not known; its site could not be read."),
 ("Proof: borrowed credibility is common",
  f'={kv("Cite a Clutch rating")}&"/"&{n}&" cite a Clutch rating; partner badges per firm (median "&Metrics!B{MED["partner"]}&")."',
  "Collect and surface third-party reviews and partner status as trust signals.",
  "Add review/badge strip near the primary CTA.", "CTA click-through before vs. after.",
  "Ratings and badges are self-reported on each firm's own site."),
]
for i, row in enumerate(plan_rows):
    for j, v in enumerate(row, 1):
        x = gp.cell(PL0 + i, j, v); x.alignment = Alignment(wrap_text=True, vertical="top")
for c, w in zip("ABCDEF", [44, 52, 52, 40, 40, 44]): gp.column_dimensions[c].width = w
gp.column_dimensions["B"].width = 52

# ---------- Dashboard
db = wb.create_sheet("Dashboard", 1)
db["A1"] = "Market Intelligence & GTM Dashboard (software engineering services)"; db["A1"].font = Font(bold=True, size=16)
db["A3"] = "Select a brand:"; db["A3"].font = BOLD
db["B3"] = BR[0]; db["B3"].font = Font(bold=True, color="0000FF")
dv = DataValidation(type="list", formula1=f"=Metrics!$A$4:$A${M1}", allow_blank=False); db.add_data_validation(dv); dv.add("B3")
def mrow(col): return f"=INDEX(Metrics!${col}$4:${col}${M1},MATCH($B$3,Metrics!$A$4:$A${M1},0))"
def bcol(c): return f"=INDEX({rng(c)},MATCH($B$3,{rng('brand')},0))"
def themes_used(sheet, S0, S1):
    return f"=SUM(INDEX('{sheet}'!$E${S0}:${L(4+NB)}${S1},0,MATCH($B$3,'{sheet}'!$E$4:${L(4+NB)}$4,0)))"
prof = [("Headline", bcol("positioning_h1")), ("Industries listed", mrow("B")), ("Engagement models named", mrow("C")),
        ("Publishes a rate", mrow("H")), ("Content engine score (0-6)", mrow("G")), ("Partner badges", mrow("E")),
        ("Positioning themes used", themes_used("Positioning", pS0, pS1)), ("CTA types used", themes_used("CTA Funnel", cS0, cS1))]
for k, (lab, f) in enumerate(prof):
    db.cell(5 + k, 1, lab).font = BOLD; db.cell(5 + k, 2, f); db.cell(5 + k, 2).alignment = Alignment(wrap_text=True, vertical="top")
db["A15"] = "Market snapshot"; db["A15"].font = Font(bold=True, size=12)
for k, lab in enumerate(["Brands compared", "Headlines AI-led", "Brands publishing a rate", "Call / consult CTA", "Whitepapers: Yes"]):
    db.cell(16 + k, 1, lab); db.cell(16 + k, 2, f"='GTM Plan'!B{KROW[lab]}")
db.column_dimensions["A"].width = 30; db.column_dimensions["B"].width = 50
def bar(title, ws, col, r0, r1, cat_col, anchor, hdr_row=None):
    c = BarChart(); c.type = "bar"; c.title = title; c.legend = None; c.height, c.width = 8, 14
    c.add_data(Reference(ws, min_col=col, min_row=hdr_row if hdr_row else r0, max_row=r1), titles_from_data=bool(hdr_row))
    c.set_categories(Reference(ws, min_col=cat_col, min_row=r0, max_row=r1)); db.add_chart(c, anchor)
bar("Content engine score (0-6)", mt, 7, 4, M1, 1, "D3", 3)
bar("Industries listed per firm", mt, 2, 4, M1, 1, "D20", 3)
bar("Brands using each positioning theme", posS, 4, pS0, pS1, 1, "D37")
bar("Brands using each CTA type", ctaS, 4, cS0, cS1, 1, "D54")

(ROOT / "excel").mkdir(exist_ok=True)
wb.save(ROOT / "excel/gtm_workbook.xlsx"); print("saved")
